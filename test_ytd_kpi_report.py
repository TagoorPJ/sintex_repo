"""YTD report regressions; no database connection or email delivery required."""
import argparse
import unittest
from datetime import date, datetime, timezone
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import load_workbook

from scripts import send_daily_kpi_report as reporting


def window(start="2026-09-10", end="2026-09-10", tz="Asia/Kolkata"):
    return reporting.build_window(argparse.Namespace(
        from_date=start, to_date=end, date=None, yesterday=False,
        days=None, timezone=tz,
    ))


def report_data():
    return {key: [] for key in (
        "daily_rows", "new_user_rows", "mobile_message_rows", "api_rows",
        "feature_usage_rows", "mobile_feature_usage_rows",
    )}


def ytd_data():
    return {
        "activity_rows": [], "new_user_rows": [],
        "engagement": {}, "mobile_user_rows": [],
    }


class YtdKpiReportTests(unittest.IsolatedAsyncioTestCase):
    def test_calendar_year_and_inclusive_timezone_boundaries(self):
        for start, end in (
            ("2026-09-10", "2026-09-10"),
            ("2025-12-20", "2026-01-01"),
            ("2024-02-01", "2024-02-29"),
        ):
            with self.subTest(end=end):
                selected = window(start, end)
                ytd = reporting.build_ytd_window(selected)
                self.assertEqual(ytd.start_day, date(selected.end_day.year, 1, 1))
                self.assertEqual(ytd.end_day, selected.end_day)
                self.assertEqual(ytd.end_utc, selected.end_utc)
                self.assertEqual(ytd.start_utc, datetime(
                    selected.end_day.year - 1, 12, 31, 18, 30, tzinfo=timezone.utc,
                ))

    async def test_ytd_loader_uses_only_historical_tables(self):
        engine = MagicMock()
        engine.connect.return_value.__aenter__ = AsyncMock(return_value=AsyncMock())
        engine.connect.return_value.__aexit__ = AsyncMock(return_value=False)
        engine.dispose = AsyncMock()
        ytd = reporting.build_ytd_window(window())
        with patch.object(reporting, "create_async_engine", return_value=engine), \
             patch.object(reporting, "table_exists", AsyncMock(return_value=True)) as exists, \
             patch.object(reporting, "table_columns", AsyncMock(return_value=set())), \
             patch.object(reporting, "fetch_all", AsyncMock(return_value=[])) as fetch:
            data = await reporting.load_ytd_report_data("postgresql+asyncpg://test", ytd)
        self.assertEqual(data, ytd_data())
        exists.assert_not_awaited()
        self.assertEqual(fetch.await_count, 4)
        for call in fetch.call_args_list:
            sql, params = call.args[1:]
            self.assertNotIn("chatbot_feature_usage_events", sql)
            self.assertNotIn("external_api_daily_counts", sql)
            for snapshot_field in ("is_registered", "stage", "status", "language"):
                self.assertNotIn(snapshot_field, sql)
            self.assertEqual(params["start_utc"], ytd.start_utc)
            self.assertEqual(params["end_utc"], ytd.end_utc)
            self.assertEqual(params["tz"], ytd.timezone_name)
            self.assertIn("created_at >= :start_utc", sql)
            self.assertIn("created_at < :end_utc", sql)
            if "FROM messages" in sql:
                self.assertIn("c.user_id", sql)
                self.assertNotIn("m.user_id", sql)
                self.assertIn("JOIN conversations c ON c.id = m.conversation_id", sql)
        mobile_sql = fetch.call_args_list[-1].args[1]
        self.assertIn("GROUP BY u.id, u.mobile_number, u.created_at", mobile_sql)
        self.assertIn("JOIN users u ON u.id = c.user_id", mobile_sql)
        self.assertIn("m.sender = 'user'", mobile_sql)
        self.assertIn("ORDER BY u.mobile_number", mobile_sql)
        self.assertIn("HAVING COUNT(*) FILTER (WHERE m.sender = 'user') > 0", mobile_sql)
        self.assertIn("COUNT(*) FILTER (WHERE m.sender = 'assistant')", mobile_sql)
        self.assertNotIn("WHERE m.sender = 'user'\n", mobile_sql)
        engine.dispose.assert_awaited_once()

    def test_workbook_uses_historical_ytd_totals_and_monthly_trends(self):
        selected = window()
        daily, ytd = report_data(), ytd_data()
        daily["daily_rows"] = [{"day": selected.end_day, "user_messages": 2}]
        ytd["activity_rows"] = [
            {"month": None, "unique_users": 2, "user_messages": 9, "assistant_messages": 11,
             "tool_messages": 1, "total_messages": 21, "active_conversations": 3, "active_days": 4},
            {"month": date(2026, 1, 1), "unique_users": 2, "user_messages": 7, "assistant_messages": 8,
             "tool_messages": 1, "total_messages": 16, "active_conversations": 2, "active_days": 3},
            {"month": date(2026, 9, 1), "unique_users": 2, "user_messages": 2, "assistant_messages": 3,
             "total_messages": 5, "active_conversations": 2, "active_days": 1},
        ]
        ytd["new_user_rows"] = [
            {"month": None, "new_users": 2}, {"month": date(2026, 1, 1), "new_users": 2},
        ]
        ytd["conversation_rows"] = [
            {"month": None, "new_conversations": 1},
            {"month": date(2026, 1, 1), "new_conversations": 1},
        ]
        ytd["engagement"] = {
            "existing_active_users": 1, "repeat_day_users": 2,
            "single_message_users": 0, "user_days": 6,
        }
        ytd["mobile_user_rows"] = [
            {"mobile_number": "+919999999999", "user_created_day": date(2025, 12, 1),
             "first_user_message_day": date(2026, 1, 1), "last_user_message_day": date(2026, 9, 10),
             "user_messages": 6, "assistant_messages": 8, "total_messages": 15,
             "active_days": 4, "active_months": 2},
            {"mobile_number": "001234567890", "user_created_day": date(2026, 1, 2),
             "first_user_message_day": date(2026, 1, 2), "last_user_message_day": date(2026, 9, 10),
             "user_messages": 3, "assistant_messages": 3, "total_messages": 6,
             "active_days": 2, "active_months": 2},
        ]
        # Extra API/feature data must never leak into the historical YTD sheet.
        ytd["api_rows"] = [{
            "day": date(2026, 1, 1), "api_name": "evolve.register",
            "total_calls": 3, "success_calls": 2, "failure_calls": 1,
        }]
        ytd["feature_usage_rows"] = [{"event_name": "RECENT FEATURE SENTINEL"}]
        filename, content = reporting.build_excel_report(daily, selected, ytd)
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(workbook.sheetnames,
                         ["Mobile Feature Usage", "Report Summary", "YTD KPI Report"])
        self.assertIn("2026-09-10", filename)
        ws = workbook["YTD KPI Report"]
        self.assertEqual(ws["C1"].value, "Period: 2026-01-01 to 2026-09-10")
        rows = list(ws.values)
        values = {row[0]: row[1] for row in rows if row[0]}
        self.assertEqual(values["Total Unique Users"], 2)
        self.assertEqual(values["New Users Came"], 2)
        self.assertEqual(values["Total User Messages"], 9)
        self.assertEqual(values["Total Agent Messages"], 11)
        self.assertNotIn("New Users Who Messaged", values)
        self.assertEqual(values["Existing Users Who Messaged"], 1)
        self.assertEqual(values["Users Active on Multiple Days"], 2)
        self.assertEqual(values["Total Messages"], 21)
        self.assertEqual(values["User Messages per Active User"], 4.5)
        self.assertEqual(values["Average Daily Active Users"], round(6 / 253, 2))
        self.assertNotIn("Flow Completion KPIs", values)
        self.assertNotIn("Evolve API Health", values)
        cells = [cell for row in rows for cell in row]
        self.assertNotIn("Feature Usage", cells)
        self.assertNotIn("RECENT FEATURE SENTINEL", cells)
        self.assertNotIn("evolve.register", cells)
        for removed_label in (
            "Total Tool Messages", "Tool Messages", "New Conversations",
            "Conversations with User Activity", "User Messages per Active Conversation",
        ):
            self.assertNotIn(removed_label, cells)
        monthly = [row for row in rows if isinstance(row[0], datetime)]
        self.assertEqual(len(monthly), 9)
        self.assertEqual(monthly[0][1:3], (2, 2))
        self.assertEqual(monthly[1][1:7], (0,) * 6)
        self.assertEqual(monthly[-1][1], 2)
        # Users recur across months; don't add monthly distinct counts.
        self.assertEqual(sum(row[1] for row in monthly), 4)
        self.assertEqual(values["Total Unique Users"], 2)
        daily_values = {row[0]: row[1] for row in workbook["Report Summary"].values if row[0]}
        self.assertEqual(daily_values["Total User Messages"], 2)
        self.assertIn("Feature Usage", daily_values)
        mobile_header = next(cell.row for cell in ws["A"] if cell.value == "Mobile Number")
        mobile_cells = [ws.cell(row=mobile_header + offset, column=1) for offset in (1, 2)]
        self.assertEqual([cell.value for cell in mobile_cells], ["+919999999999", "001234567890"])
        self.assertTrue(all(cell.data_type == "s" and cell.number_format == "@" for cell in mobile_cells))
        analytics = [list(row) for row in ws.iter_rows(
            min_row=mobile_header + 1, max_row=mobile_header + 2, values_only=True,
        )]
        self.assertEqual(analytics[0][1:4], [datetime(2025, 12, 1), datetime(2026, 1, 1), datetime(2026, 9, 10)])
        self.assertEqual(analytics[0][4:], [6, 8, 15, 4, 2, 1.5, 0])
        self.assertEqual(analytics[1][4:], [3, 3, 6, 2, 2, 1.5, 0])
        self.assertEqual(sum(row[4] for row in analytics), values["Total User Messages"])
        self.assertEqual(sum(row[5] for row in analytics), values["Total Agent Messages"])
        self.assertEqual(sum(row[6] for row in analytics), values["Total Messages"])

    def test_empty_year_and_undefined_ratios(self):
        ytd_window = reporting.build_ytd_window(window("2024-02-29", "2024-02-29"))
        tables = reporting.build_ytd_summary_tables(ytd_data(), ytd_window)
        summary = {row[0]: row[1] for row in tables[0][2]}
        self.assertEqual(summary["Total Unique Users"], 0)
        self.assertEqual(summary["User Messages per Active User"], "N/A")
        self.assertEqual(summary["Average Daily Active Users"], 0)
        self.assertEqual([row[0] for row in tables[1][2]], [date(2024, 1, 1), date(2024, 2, 1)])
        self.assertTrue(all(row[1:] == [0] * 6 for row in tables[1][2]))
        self.assertEqual(tables[2][0], "YTD Mobile Analytics")
        self.assertEqual(tables[2][2], [])

    def test_mobile_recency_uses_report_end_date_and_handles_single_day(self):
        data = ytd_data()
        data["mobile_user_rows"] = [{
            "mobile_number": "001234567890", "user_created_day": date(2023, 12, 31),
            "first_user_message_day": date(2024, 2, 28), "last_user_message_day": date(2024, 2, 28),
            "user_messages": 1, "assistant_messages": 0, "total_messages": 1,
            "active_days": 1, "active_months": 1,
        }]
        ytd_window = reporting.build_ytd_window(window("2024-02-29", "2024-02-29"))
        mobile_table = reporting.build_ytd_summary_tables(data, ytd_window)[2]
        self.assertEqual(mobile_table[2][0][4:], [1, 0, 1, 1, 1, 1.0, 1])

    async def test_main_loads_ytd_data_for_excel_without_sending_in_dry_run(self):
        args = argparse.Namespace(
            from_date=None, to_date=None, date="2026-09-10", yesterday=False,
            days=None, timezone="Asia/Kolkata", log_file=None, dry_run=True,
            to=None, excel_output=None, subject_prefix="Test",
        )
        daily, ytd = report_data(), ytd_data()
        with patch.object(reporting, "parse_args", return_value=args), \
             patch.object(reporting, "setup_logging"), \
             patch.dict(reporting.os.environ, {"DATABASE_URL": "postgresql+asyncpg://test"}), \
             patch.object(reporting, "load_report_data", AsyncMock(return_value=daily)) as load, \
             patch.object(reporting, "load_ytd_report_data", AsyncMock(return_value=ytd)) as load_ytd, \
             patch.object(reporting, "build_excel_report", return_value=("test.xlsx", b"test")) as build, \
             patch.object(reporting, "send_email") as send, patch("builtins.print"):
            await reporting.async_main()
        load.assert_awaited_once()
        load_ytd.assert_awaited_once()
        self.assertEqual(load.call_args.args[1].start_day, date(2026, 9, 10))
        self.assertEqual(load_ytd.call_args.args[1].start_day, date(2026, 1, 1))
        self.assertEqual(load_ytd.call_args.args[1].end_day, date(2026, 9, 10))
        self.assertIs(build.call_args.args[2], ytd)
        send.assert_not_called()


if __name__ == "__main__":
    unittest.main()
