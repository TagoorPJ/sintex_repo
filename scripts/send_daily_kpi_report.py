from __future__ import annotations

import argparse
import asyncio
import logging
import os
import smtplib
import sys
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from email.message import EmailMessage
from html import escape
from io import BytesIO
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from dotenv import load_dotenv
from sqlalchemy import inspect, text
from sqlalchemy.ext.asyncio import create_async_engine


def find_project_root() -> Path:
    script_path = Path(__file__).resolve()
    for candidate in (script_path.parent, *script_path.parents):
        if (candidate / "alembic.ini").exists() and (candidate / "app").is_dir():
            return candidate
    for candidate in (script_path.parent, *script_path.parents):
        if (candidate / ".env").exists():
            return candidate
    return script_path.parents[1]


BASE_DIR = find_project_root()
load_dotenv(BASE_DIR / ".env")

DEFAULT_TIMEZONE = "Asia/Kolkata"
DEFAULT_SUBJECT_PREFIX = "Sintex Agent KPI Report"
logger = logging.getLogger("daily_kpi_report")


def setup_logging(log_file: str | None = None) -> None:
    handlers: list[logging.Handler] = [logging.StreamHandler(sys.stdout)]
    if log_file:
        path = Path(log_file).expanduser()
        path.parent.mkdir(parents=True, exist_ok=True)
        handlers.append(logging.FileHandler(path, encoding="utf-8"))

    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s | %(levelname)s | %(name)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S %Z",
        handlers=handlers,
        force=True,
    )

RESPONSIVE_EMAIL_CSS = """
<style>
  @media only screen and (max-width: 640px) {
    .page {
      width: 100% !important;
      max-width: 100% !important;
      padding: 12px !important;
      box-sizing: border-box !important;
    }
    .hero {
      padding: 18px 16px !important;
      border-radius: 8px !important;
    }
    .hero-title {
      font-size: 22px !important;
      line-height: 1.25 !important;
    }
    .hero-copy {
      font-size: 12px !important;
      line-height: 1.45 !important;
    }
    .kpi-table,
    .kpi-table tbody,
    .kpi-table tr {
      display: block !important;
      width: 100% !important;
    }
    .kpi-cell {
      display: block !important;
      width: 100% !important;
      padding: 6px 0 !important;
      box-sizing: border-box !important;
    }
    .kpi-card {
      min-height: 0 !important;
      padding: 14px !important;
    }
    .section {
      padding: 14px !important;
      border-radius: 8px !important;
    }
    .desktop-table {
      display: none !important;
      max-height: 0 !important;
      overflow: hidden !important;
    }
    .mobile-records {
      display: block !important;
      max-height: none !important;
      overflow: visible !important;
    }
  }
</style>
"""

FLOW_COMPLETION_KPIS = {
    "evolve.register": {
        "name": "Completed Registrations",
        "flow": "Registration",
        "signal": "register API success response",
    },
    "evolve.balance_inquiry": {
        "name": "Reward Balance Checks",
        "flow": "My Rewards",
        "signal": "balance inquiry API success response",
    },
    "evolve.points_history": {
        "name": "Recent Transaction Views",
        "flow": "My Rewards",
        "signal": "points history API success response",
    },
    "evolve.tier_status": {
        "name": "Membership Tier Checks",
        "flow": "My Rewards",
        "signal": "tier status API success response",
    },
    "evolve.reward_statement": {
        "name": "Reward Statement Views",
        "flow": "My Rewards",
        "signal": "reward statement API success response",
    },
    "evolve.find_dealer": {
        "name": "Dealer List Loads",
        "flow": "Upload Purchase",
        "signal": "find dealer API success response",
    },
    "evolve.capture_qr": {
        "name": "QR Uploads Completed",
        "flow": "Upload Purchase",
        "signal": "capture QR API success or HTTP 200 response",
    },
    "evolve.claim_create": {
        "name": "Claims Created",
        "flow": "Upload Purchase",
        "signal": "claim create API success response",
    },
    "evolve.claim_status": {
        "name": "Claim Status Checks",
        "flow": "Upload Purchase",
        "signal": "claim status API success response",
    },
    "evolve.support_ticket": {
        "name": "Support Tickets Created",
        "flow": "Support",
        "signal": "support ticket API success response",
    },
}

# Share the exact completion categories used by the application.
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))
from app.analytics.feature_events import (
    FEATURE_EVENTS,
    FEATURE_RESULT_METHODS,
)

FEATURE_USAGE_COLUMNS = [name for name, _ in FEATURE_EVENTS.values()]


@dataclass
class DateWindow:
    start_day: date
    end_day: date
    timezone_name: str
    start_utc: datetime
    end_utc: datetime


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Email daily Sintex Agent KPI report."
    )
    parser.add_argument(
        "--date",
        help="Generate the report for one date in YYYY-MM-DD format.",
    )
    parser.add_argument(
        "--from-date",
        "--start-date",
        dest="from_date",
        help="First date of an inclusive report range in YYYY-MM-DD format.",
    )
    parser.add_argument(
        "--to-date",
        "--end-date",
        dest="to_date",
        help="Last date of an inclusive report range in YYYY-MM-DD format.",
    )
    parser.add_argument(
        "--yesterday",
        action="store_true",
        help="Use yesterday in --timezone as the report end date.",
    )
    parser.add_argument(
        "--days",
        type=int,
        default=None,
        help="Rolling number of days ending at --date or today. Default: 1.",
    )
    parser.add_argument(
        "--timezone",
        default=os.getenv("KPI_REPORT_TIMEZONE", DEFAULT_TIMEZONE),
        help=f"Business timezone. Default: {DEFAULT_TIMEZONE}.",
    )
    parser.add_argument(
        "--to",
        help=(
            "Comma-separated recipients. Defaults to KPI_REPORT_TO, "
            "or KPI_EMAIL_TO."
        ),
    )
    parser.add_argument(
        "--subject-prefix",
        default=os.getenv("KPI_REPORT_SUBJECT_PREFIX", DEFAULT_SUBJECT_PREFIX),
        help=f"Email subject prefix. Default: {DEFAULT_SUBJECT_PREFIX}.",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Print report summary without sending email.",
    )
    parser.add_argument(
        "--log-file",
        default=os.getenv("KPI_REPORT_LOG_FILE"),
        help="Optional file path for timestamped script logs.",
    )
    parser.add_argument(
        "--excel-output",
        default=os.getenv("KPI_REPORT_EXCEL_OUTPUT"),
        help="Optional path to save the generated Excel report.",
    )
    return parser.parse_args()


def split_csv(value: str | None) -> list[str]:
    if not value:
        return []
    return [item.strip() for item in value.split(",") if item.strip()]


def resolve_timezone(timezone_name: str):
    try:
        return ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError:
        normalized = (timezone_name or "").strip().lower()
        if normalized in {"asia/kolkata", "asia/calcutta", "ist"}:
            return timezone(timedelta(hours=5, minutes=30), name="Asia/Kolkata")
        if normalized in {"utc", "etc/utc"}:
            return timezone.utc
        raise


def build_window(args: argparse.Namespace) -> DateWindow:
    days = getattr(args, "days", None)
    from_date = getattr(args, "from_date", None)
    to_date = getattr(args, "to_date", None)
    range_selected = bool(from_date or to_date)

    if days is not None and days < 1:
        raise ValueError("--days must be at least 1")
    if bool(from_date) != bool(to_date):
        raise ValueError("Use --from-date and --to-date together")
    if range_selected and (args.date or args.yesterday or days is not None):
        raise ValueError(
            "A date range cannot be combined with --date, --yesterday, or --days"
        )
    if args.date and args.yesterday:
        raise ValueError("Use either --date or --yesterday, not both")

    tz = resolve_timezone(args.timezone)
    if range_selected:
        start_day = datetime.strptime(from_date, "%Y-%m-%d").date()
        end_day = datetime.strptime(to_date, "%Y-%m-%d").date()
        if start_day > end_day:
            raise ValueError("--from-date cannot be later than --to-date")
    elif args.date:
        end_day = datetime.strptime(args.date, "%Y-%m-%d").date()
        start_day = end_day - timedelta(days=(days or 1) - 1)
    elif args.yesterday:
        end_day = datetime.now(tz).date() - timedelta(days=1)
        start_day = end_day - timedelta(days=(days or 1) - 1)
    else:
        end_day = datetime.now(tz).date()
        start_day = end_day - timedelta(days=(days or 1) - 1)

    local_start = datetime.combine(start_day, time.min, tzinfo=tz)
    local_end = datetime.combine(end_day + timedelta(days=1), time.min, tzinfo=tz)

    return DateWindow(
        start_day=start_day,
        end_day=end_day,
        timezone_name=args.timezone,
        start_utc=local_start.astimezone(timezone.utc),
        end_utc=local_end.astimezone(timezone.utc),
    )


def build_ytd_window(window: DateWindow) -> DateWindow:
    """Calendar year to date, ending on the selected report's final day."""
    start_day = date(window.end_day.year, 1, 1)
    local_start = datetime.combine(
        start_day, time.min, tzinfo=resolve_timezone(window.timezone_name)
    )
    return DateWindow(
        start_day=start_day,
        end_day=window.end_day,
        timezone_name=window.timezone_name,
        start_utc=local_start.astimezone(timezone.utc),
        end_utc=window.end_utc,
    )


def day_range(window: DateWindow) -> list[date]:
    days = []
    current = window.start_day
    while current <= window.end_day:
        days.append(current)
        current += timedelta(days=1)
    return days


async def table_exists(conn, table_name: str) -> bool:
    return await conn.run_sync(
        lambda sync_conn: inspect(sync_conn).has_table(table_name)
    )


async def table_columns(conn, table_name: str) -> set[str]:
    return await conn.run_sync(
        lambda sync_conn: {
            column["name"]
            for column in inspect(sync_conn).get_columns(table_name)
        }
    )


async def fetch_all(conn, statement: str, params: dict[str, Any]) -> list[dict[str, Any]]:
    result = await conn.execute(text(statement), params)
    return [dict(row._mapping) for row in result.fetchall()]


def _api_int(row: dict[str, Any], key: str) -> int:
    return int(row.get(key) or 0)


def aggregate_flow_completion_rows(api_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    records: dict[tuple[date, str], dict[str, Any]] = {}
    for row in api_rows:
        api_name = str(row.get("api_name") or "")
        definition = FLOW_COMPLETION_KPIS.get(api_name)
        if not definition:
            continue

        key = (row["day"], api_name)
        record = records.setdefault(
            key,
            {
                "day": row["day"],
                "api_name": api_name,
                "kpi_name": definition["name"],
                "flow": definition["flow"],
                "signal": definition["signal"],
                "total_calls": 0,
                "success_calls": 0,
                "failure_calls": 0,
            },
        )
        record["total_calls"] += _api_int(row, "total_calls")
        record["success_calls"] += _api_int(row, "success_calls")
        record["failure_calls"] += _api_int(row, "failure_calls")

    return sorted(
        records.values(),
        key=lambda item: (item["day"], item["flow"], item["kpi_name"]),
    )


async def load_report_data(
    database_url: str, window: DateWindow, *, include_feature_usage: bool = True
) -> dict[str, Any]:
    engine = create_async_engine(database_url, pool_pre_ping=True)
    try:
        async with engine.connect() as conn:
            daily_rows = await fetch_all(
                conn,
                """
                SELECT
                    (created_at AT TIME ZONE :tz)::date AS day,
                    COUNT(DISTINCT user_id) FILTER (WHERE sender = 'user') AS unique_users,
                    COUNT(*) FILTER (WHERE sender = 'user') AS user_messages,
                    COUNT(*) FILTER (WHERE sender = 'assistant') AS assistant_messages,
                    COUNT(*) AS total_messages
                FROM messages
                WHERE created_at >= :start_utc
                  AND created_at < :end_utc
                GROUP BY day
                ORDER BY day
                """,
                {
                    "tz": window.timezone_name,
                    "start_utc": window.start_utc,
                    "end_utc": window.end_utc,
                },
            )

            new_user_rows = await fetch_all(
                conn,
                """
                SELECT
                    (created_at AT TIME ZONE :tz)::date AS day,
                    COUNT(*) AS new_users,
                    COUNT(*) FILTER (
                        WHERE user_id_evolve IS NOT NULL
                          AND user_id_evolve NOT IN ('', 'null', 'none', 'undefined')
                    ) AS new_users_with_evolve_id
                FROM users
                WHERE created_at >= :start_utc
                  AND created_at < :end_utc
                GROUP BY day
                ORDER BY day
                """,
                {
                    "tz": window.timezone_name,
                    "start_utc": window.start_utc,
                    "end_utc": window.end_utc,
                },
            )

            mobile_message_rows = await fetch_all(
                conn,
                """
                SELECT
                    u.mobile_number,
                    COUNT(*) AS total_messages
                FROM messages m
                JOIN users u ON u.id = m.user_id
                WHERE m.sender = 'user'
                  AND m.created_at >= :start_utc
                  AND m.created_at < :end_utc
                GROUP BY u.mobile_number
                ORDER BY u.mobile_number
                """,
                {
                    "start_utc": window.start_utc,
                    "end_utc": window.end_utc,
                },
            )

            api_rows: list[dict[str, Any]] = []
            if await table_exists(conn, "external_api_daily_counts"):
                api_columns = await table_columns(conn, "external_api_daily_counts")
                status_codes_select = (
                    "status_codes"
                    if "status_codes" in api_columns
                    else "NULL::text AS status_codes"
                )
                error_messages_select = (
                    "error_messages"
                    if "error_messages" in api_columns
                    else "NULL::text AS error_messages"
                )
                api_rows = await fetch_all(
                    conn,
                    f"""
                    SELECT
                        call_date AS day,
                        api_name,
                        method,
                        endpoint,
                        total_calls,
                        success_calls,
                        failure_calls,
                        total_duration_ms,
                        max_duration_ms,
                        last_status_code,
                        last_error,
                        {status_codes_select},
                        {error_messages_select}
                    FROM external_api_daily_counts
                    WHERE call_date >= :start_day
                      AND call_date <= :end_day
                    ORDER BY call_date, api_name, method, endpoint
                    """,
                    {
                        "start_day": window.start_day,
                        "end_day": window.end_day,
                    },
                )

            feature_usage_rows: list[dict[str, Any]] = []
            mobile_feature_usage_rows: list[dict[str, Any]] = []
            if include_feature_usage and await table_exists(conn, "chatbot_feature_usage_events"):
                feature_usage_rows = await fetch_all(
                    conn,
                    """
                    SELECT
                        event_date AS day,
                        event_key,
                        event_name,
                        flow_name,
                        COUNT(DISTINCT user_id) AS unique_users,
                        COUNT(*) AS usage_count,
                        STRING_AGG(DISTINCT detection_method, ', ' ORDER BY detection_method) AS detection_methods
                    FROM chatbot_feature_usage_events
                    WHERE detection_method = ANY(:result_methods)
                      AND event_key = ANY(:feature_keys)
                      AND event_date >= :start_day
                      AND event_date <= :end_day
                    GROUP BY event_date, event_key, event_name, flow_name
                    ORDER BY event_date, flow_name, event_name
                    """,
                    {
                        "result_methods": list(FEATURE_RESULT_METHODS),
                        "feature_keys": list(FEATURE_EVENTS),
                        "start_day": window.start_day,
                        "end_day": window.end_day,
                    },
                )
                mobile_feature_usage_rows = await fetch_all(
                    conn,
                    """
                    SELECT
                        mobile_number,
                        event_name,
                        COUNT(*) AS usage_count
                    FROM chatbot_feature_usage_events
                    WHERE detection_method = ANY(:result_methods)
                      AND event_key = ANY(:feature_keys)
                      AND event_date >= :start_day
                      AND event_date <= :end_day
                    GROUP BY mobile_number, event_name
                    ORDER BY mobile_number, event_name
                    """,
                    {
                        "result_methods": list(FEATURE_RESULT_METHODS),
                        "feature_keys": list(FEATURE_EVENTS),
                        "start_day": window.start_day,
                        "end_day": window.end_day,
                    },
                )

    finally:
        await engine.dispose()

    return {
        "daily_rows": daily_rows,
        "new_user_rows": new_user_rows,
        "mobile_message_rows": mobile_message_rows,
        "api_rows": api_rows,
        "feature_usage_rows": feature_usage_rows,
        "mobile_feature_usage_rows": mobile_feature_usage_rows,
    }


async def load_ytd_report_data(database_url: str, window: DateWindow) -> dict[str, Any]:
    """Historical activity only: no API counters, feature events or mutable state."""
    params = {
        "tz": window.timezone_name,
        "start_utc": window.start_utc,
        "end_utc": window.end_utc,
    }
    engine = create_async_engine(database_url, pool_pre_ping=True)
    try:
        async with engine.connect() as conn:
            # The conversation owns the user relationship, including messages
            # saved before messages.user_id was introduced/backfilled.
            activity_rows = await fetch_all(conn, """
                WITH activity AS (
                    SELECT c.user_id, m.sender,
                           (m.created_at AT TIME ZONE :tz)::date AS day,
                           DATE_TRUNC('month', m.created_at AT TIME ZONE :tz)::date AS month
                    FROM messages m
                    JOIN conversations c ON c.id = m.conversation_id
                    WHERE m.created_at >= :start_utc AND m.created_at < :end_utc
                )
                SELECT month,
                       COUNT(DISTINCT user_id) FILTER (WHERE sender = 'user') AS unique_users,
                       COUNT(*) FILTER (WHERE sender = 'user') AS user_messages,
                       COUNT(*) FILTER (WHERE sender = 'assistant') AS assistant_messages,
                       COUNT(*) AS total_messages,
                       COUNT(DISTINCT day) FILTER (WHERE sender = 'user') AS active_days
                FROM activity
                GROUP BY GROUPING SETS ((), (month))
                ORDER BY month NULLS FIRST
            """, params)
            new_user_rows = await fetch_all(conn, """
                SELECT DATE_TRUNC('month', created_at AT TIME ZONE :tz)::date AS month,
                       COUNT(*) AS new_users
                FROM users
                WHERE created_at >= :start_utc AND created_at < :end_utc
                GROUP BY GROUPING SETS ((), (month))
                ORDER BY month NULLS FIRST
            """, params)
            engagement_rows = await fetch_all(conn, """
                WITH user_activity AS (
                    SELECT c.user_id,
                           COUNT(*) AS user_messages,
                           COUNT(DISTINCT (m.created_at AT TIME ZONE :tz)::date) AS active_days
                    FROM messages m
                    JOIN conversations c ON c.id = m.conversation_id
                    WHERE m.sender = 'user'
                      AND m.created_at >= :start_utc AND m.created_at < :end_utc
                    GROUP BY c.user_id
                )
                SELECT COUNT(*) FILTER (WHERE u.created_at < :start_utc) AS existing_active_users,
                       COUNT(*) FILTER (WHERE a.active_days >= 2) AS repeat_day_users,
                       COUNT(*) FILTER (WHERE a.user_messages = 1) AS single_message_users,
                       COALESCE(SUM(a.active_days), 0) AS user_days
                FROM user_activity a
                JOIN users u ON u.id = a.user_id
            """, params)
            mobile_user_rows = await fetch_all(conn, """
                SELECT u.mobile_number,
                       (u.created_at AT TIME ZONE :tz)::date AS user_created_day,
                       MIN((m.created_at AT TIME ZONE :tz)::date)
                           FILTER (WHERE m.sender = 'user') AS first_user_message_day,
                       MAX((m.created_at AT TIME ZONE :tz)::date)
                           FILTER (WHERE m.sender = 'user') AS last_user_message_day,
                       COUNT(*) FILTER (WHERE m.sender = 'user') AS user_messages,
                       COUNT(*) FILTER (WHERE m.sender = 'assistant') AS assistant_messages,
                       COUNT(*) AS total_messages,
                       COUNT(DISTINCT (m.created_at AT TIME ZONE :tz)::date)
                           FILTER (WHERE m.sender = 'user') AS active_days,
                       COUNT(DISTINCT DATE_TRUNC('month', m.created_at AT TIME ZONE :tz))
                           FILTER (WHERE m.sender = 'user') AS active_months
                FROM messages m
                JOIN conversations c ON c.id = m.conversation_id
                JOIN users u ON u.id = c.user_id
                WHERE m.created_at >= :start_utc AND m.created_at < :end_utc
                GROUP BY u.id, u.mobile_number, u.created_at
                HAVING COUNT(*) FILTER (WHERE m.sender = 'user') > 0
                ORDER BY u.mobile_number
            """, params)
    finally:
        await engine.dispose()
    return {
        "activity_rows": activity_rows,
        "new_user_rows": new_user_rows,
        "engagement": engagement_rows[0] if engagement_rows else {},
        "mobile_user_rows": mobile_user_rows,
    }


def build_ytd_summary_tables(
    report: dict[str, Any], window: DateWindow
) -> list[tuple[str, list[str], list[list[Any]]]]:
    activity = {row["month"]: row for row in report["activity_rows"]}
    users = {row["month"]: row for row in report["new_user_rows"]}
    total = activity.get(None, {})
    engagement = report["engagement"]

    def count(row, key):
        return int(row.get(key) or 0)

    def ratio(numerator, denominator):
        return round(numerator / denominator, 2) if denominator else "N/A"

    unique_users = count(total, "unique_users")
    user_messages = count(total, "user_messages")
    new_users = count(users.get(None, {}), "new_users")
    calendar_days = (window.end_day - window.start_day).days + 1
    summary = [
        ["Total Unique Users", unique_users, "Distinct users who sent a user message during YTD; counted once across the period."],
        ["New Users Came", new_users, "User records created during YTD; does not imply Evolve registration."],
        ["Existing Users Who Messaged", count(engagement, "existing_active_users"), "Users created before January 1 who sent a user message during YTD."],
        ["Users Active on Multiple Days", count(engagement, "repeat_day_users"), "Users who sent messages on at least two distinct local dates during YTD."],
        ["Single-Message Users", count(engagement, "single_message_users"), "Users with exactly one user message across YTD."],
        ["Total User Messages", user_messages, "Stored messages with sender=user during YTD."],
        ["Total Agent Messages", count(total, "assistant_messages"), "Stored messages with sender=assistant during YTD."],
        ["Total Messages", count(total, "total_messages"), "All stored message senders during YTD."],
        ["Days with User Activity", count(total, "active_days"), "Distinct local dates with at least one user message."],
        ["Average Daily Active Users", ratio(count(engagement, "user_days"), calendar_days), "Sum of daily distinct messaging users / all calendar days in the selected YTD period."],
        ["User Messages per Active User", ratio(user_messages, unique_users), "YTD user messages / YTD distinct messaging users."],
    ]
    monthly = []
    month = window.start_day.replace(day=1)
    while month <= window.end_day:
        row = activity.get(month, {})
        monthly.append([
            month, count(row, "unique_users"),
            count(users.get(month, {}), "new_users"),
            count(row, "user_messages"), count(row, "assistant_messages"),
            count(row, "total_messages"), count(row, "active_days"),
        ])
        month = (month.replace(day=28) + timedelta(days=4)).replace(day=1)
    mobile_analytics = []
    for row in report["mobile_user_rows"]:
        last_day = row.get("last_user_message_day")
        mobile_analytics.append([
            str(row["mobile_number"]), row.get("user_created_day"),
            row.get("first_user_message_day"), last_day,
            count(row, "user_messages"), count(row, "assistant_messages"),
            count(row, "total_messages"), count(row, "active_days"),
            count(row, "active_months"),
            ratio(count(row, "user_messages"), count(row, "active_days")),
            (window.end_day - last_day).days if last_day is not None else None,
        ])
    return [
        ("YTD KPI Summary", ["KPI", "Value", "Definition"], summary),
        ("Monthly Activity", [
            "Month", "Unique Messaging Users", "New Users", "User Messages",
            "Agent Messages", "Total Messages", "Days with User Activity",
        ], monthly),
        ("YTD Mobile Analytics", [
            "Mobile Number", "User Record Created", "First YTD User Message",
            "Last YTD User Message", "User Messages", "Agent Messages",
            "Total Messages", "Active Days", "Active Months",
            "User Messages per Active Day", "Days Since Last User Message",
        ], mobile_analytics),
    ]


def index_by_day(rows: list[dict[str, Any]]) -> dict[date, dict[str, Any]]:
    return {row["day"]: row for row in rows}


def fmt_int(value: Any) -> str:
    return f"{int(value or 0):,}"


def fmt_ms(value: Any) -> str:
    return f"{float(value or 0):,.2f} ms"


def pct(numerator: int, denominator: int) -> str:
    if denominator <= 0:
        return "0.0%"
    return f"{(numerator / denominator) * 100:.1f}%"


def average_ms(total_duration_ms: Any, total_calls: Any) -> float:
    calls = int(total_calls or 0)
    if calls <= 0:
        return 0.0
    return float(total_duration_ms or 0.0) / calls


def format_status_codes(row: dict[str, Any]) -> str:
    values = set()
    raw_codes = str(row.get("status_codes") or "")
    for item in raw_codes.replace("\n", ",").split(","):
        code = item.strip()
        if code and code.lower() not in {"none", "null"}:
            values.add(code)
    last_status_code = row.get("last_status_code")
    if last_status_code is not None:
        values.add(str(last_status_code))
    if not values:
        return "-"
    return ", ".join(sorted(values, key=lambda value: (not value.isdigit(), value)))


def format_error_messages(row: dict[str, Any]) -> str:
    raw_errors = row.get("error_messages") or row.get("last_error") or ""
    seen = set()
    errors = []
    for item in str(raw_errors).split("\n---\n"):
        message = " ".join(item.split())
        if not message or message in seen:
            continue
        seen.add(message)
        errors.append(message[:220])
    if not errors:
        return "-"
    suffix = "" if len(errors) <= 3 else f" (+{len(errors) - 3} more)"
    return "; ".join(errors[:3]) + suffix


def td(value: Any, *, align: str = "left") -> str:
    return (
        f'<td style="padding:10px;border-bottom:1px solid #e5e7eb;'
        f'text-align:{align};vertical-align:top;line-height:1.35;'
        f'word-break:break-word;">{escape(str(value))}</td>'
    )


def th(value: str, *, align: str = "left") -> str:
    return (
        f'<th style="padding:10px;border-bottom:1px solid #d1d5db;'
        f'text-align:{align};font-size:12px;color:#374151;'
        f'vertical-align:bottom;line-height:1.35;">{escape(value)}</th>'
    )


def render_mobile_records(headers: list[tuple[str, str]], rows: list[list[Any]]) -> str:
    cards = []
    for row in rows:
        pairs = []
        for value, (label, align) in zip(row, headers):
            value_align = "right" if align == "right" else "left"
            pairs.append(
                "<tr>"
                '<td style="padding:7px 0;color:#6b7280;font-size:12px;'
                'line-height:1.35;vertical-align:top;width:42%;">'
                f"{escape(label)}"
                "</td>"
                f'<td style="padding:7px 0;color:#111827;font-size:13px;'
                f'font-weight:600;line-height:1.35;text-align:{value_align};'
                f'vertical-align:top;word-break:break-word;">'
                f"{escape(str(value))}"
                "</td>"
                "</tr>"
            )
        cards.append(
            '<div style="border:1px solid #e5e7eb;border-radius:8px;'
            'padding:12px;margin-top:10px;background:#ffffff;">'
            '<table role="presentation" style="width:100%;border-collapse:collapse;">'
            f"{''.join(pairs)}"
            "</table>"
            "</div>"
        )
    return (
        '<div class="mobile-records" style="display:none;mso-hide:all;">'
        f"{''.join(cards)}"
        "</div>"
    )


def render_table(headers: list[tuple[str, str]], rows: list[list[Any]]) -> str:
    if not rows:
        return '<p style="color:#6b7280;margin:8px 0 0;">No records for this period.</p>'
    header_html = "".join(th(label, align=align) for label, align in headers)
    body_html = []
    for row in rows:
        cells = []
        for value, (_, align) in zip(row, headers):
            cells.append(td(value, align=align))
        body_html.append("<tr>" + "".join(cells) + "</tr>")
    min_width = max(520, len(headers) * 110)
    return (
        '<div class="desktop-table" style="width:100%;overflow-x:auto;'
        '-webkit-overflow-scrolling:touch;">'
        f'<table style="width:100%;min-width:{min_width}px;'
        'border-collapse:collapse;font-size:13px;">'
        f"<thead><tr>{header_html}</tr></thead>"
        f"<tbody>{''.join(body_html)}</tbody>"
        "</table>"
        "</div>"
        f"{render_mobile_records(headers, rows)}"
    )


def build_mobile_feature_usage_sheet(
    report: dict[str, Any],
) -> tuple[list[str], list[list[Any]]]:
    mobile_records: dict[str, dict[str, int]] = {}
    total_messages: dict[str, int] = {}
    extra_features = set()

    for row in report["mobile_message_rows"]:
        mobile = str(row.get("mobile_number") or "").strip() or "(unknown)"
        mobile_records.setdefault(mobile, {})
        total_messages[mobile] = total_messages.get(mobile, 0) + int(
            row.get("total_messages") or 0
        )

    for row in report["mobile_feature_usage_rows"]:
        mobile = str(row.get("mobile_number") or "").strip() or "(unknown)"
        event_name = str(row.get("event_name") or "").strip()
        if not event_name:
            continue
        mobile_records.setdefault(mobile, {})
        mobile_records[mobile][event_name] = mobile_records[mobile].get(
            event_name, 0
        ) + int(row.get("usage_count") or 0)
        if event_name not in FEATURE_USAGE_COLUMNS:
            extra_features.add(event_name)

    feature_columns = FEATURE_USAGE_COLUMNS + sorted(extra_features)
    headers = ["Mobile Number", *feature_columns, "Total Messages"]
    rows = []
    for mobile in sorted(mobile_records):
        feature_counts = mobile_records[mobile]
        rows.append([
            mobile,
            *[feature_counts.get(feature, 0) for feature in feature_columns],
            total_messages.get(mobile, 0),
        ])

    return headers, rows


def build_top_kpi_totals(report: dict[str, Any], window: DateWindow) -> dict[str, int]:
    daily_by_day = index_by_day(report["daily_rows"])
    new_users_by_day = index_by_day(report["new_user_rows"])
    unique_mobiles = {
        str(row.get("mobile_number") or "").strip()
        for row in report["mobile_message_rows"]
        if str(row.get("mobile_number") or "").strip()
    }

    totals = {
        "unique_users": len(unique_mobiles),
        "user_messages": 0,
        "assistant_messages": 0,
        "new_users": 0,
    }

    for day in day_range(window):
        daily = daily_by_day.get(day, {})
        new_users = new_users_by_day.get(day, {})
        totals["user_messages"] += int(daily.get("user_messages") or 0)
        totals["assistant_messages"] += int(daily.get("assistant_messages") or 0)
        totals["new_users"] += int(new_users.get("new_users") or 0)

    return totals


def build_report_summary_tables(
    report: dict[str, Any],
    window: DateWindow,
    *,
    include_feature_usage: bool = True,
) -> list[tuple[str, list[str], list[list[Any]]]]:
    api_rows = report["api_rows"]
    feature_usage_rows = report["feature_usage_rows"] if include_feature_usage else []
    completion_rows = aggregate_flow_completion_rows(api_rows)
    top_totals = build_top_kpi_totals(report, window)
    api_totals = {
        "completed_flow_actions": 0,
        "api_calls": 0,
        "api_success": 0,
        "api_failure": 0,
    }

    for row in api_rows:
        api_totals["api_calls"] += int(row.get("total_calls") or 0)
        api_totals["api_success"] += int(row.get("success_calls") or 0)
        api_totals["api_failure"] += int(row.get("failure_calls") or 0)

    for row in completion_rows:
        api_totals["completed_flow_actions"] += int(row.get("success_calls") or 0)

    top_summary = [
        ["Total Unique Users", int(top_totals["unique_users"])],
        ["New Users Came", int(top_totals["new_users"])],
        ["Total User Messages", int(top_totals["user_messages"])],
        ["Total Agent Messages", int(top_totals["assistant_messages"])],
    ]

    feature_usage = []
    for row in feature_usage_rows:
        feature_usage.append([
            row["day"].isoformat(),
            row.get("event_name") or "",
            row.get("flow_name") or "",
            int(row.get("unique_users") or 0),
            int(row.get("usage_count") or 0),
        ])

    completion = []
    for row in completion_rows:
        completion.append([
            row["day"].isoformat(),
            row.get("kpi_name") or "",
            row.get("flow") or "",
            row.get("api_name") or "",
            int(row.get("success_calls") or 0),
            int(row.get("total_calls") or 0),
            int(row.get("failure_calls") or 0),
        ])

    api_health = []
    for row in api_rows:
        total = int(row.get("total_calls") or 0)
        api_health.append([
            row["day"].isoformat(),
            row.get("api_name") or "",
            row.get("method") or "",
            row.get("endpoint") or "",
            total,
            int(row.get("success_calls") or 0),
            int(row.get("failure_calls") or 0),
            round(average_ms(row.get("total_duration_ms"), total), 2),
            round(float(row.get("max_duration_ms") or 0.0), 2),
            format_status_codes(row),
            format_error_messages(row),
        ])

    return [
        ("Top KPI Summary", ["KPI", "Value"], top_summary),
        *([(
            "Feature Usage",
            ["Day", "Feature", "Flow", "Unique Users", "Usage Count"],
            feature_usage,
        )] if include_feature_usage else []),
        (
            "Flow Completion KPIs",
            ["Day", "KPI", "Flow", "Completion API", "Completed", "Attempts", "Failed"],
            completion,
        ),
        (
            "Evolve API Health",
            [
                "Day",
                "API",
                "Method",
                "Endpoint",
                "Total Hits",
                "Successful Hits",
                "Failed Hits",
                "Average Time MS",
                "Max Time MS",
                "Status Codes",
                "Errors",
            ],
            api_health,
        ),
    ]


def _set_column_widths(ws) -> None:
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        ws.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 45)


def _append_excel_table(ws, title: str, headers: list[str], rows: list[list[Any]], start_row: int) -> int:
    from openpyxl.styles import Alignment, Font, PatternFill

    title_cell = ws.cell(row=start_row, column=1, value=title)
    title_cell.font = Font(bold=True, size=13, color="111827")
    start_row += 1

    header_fill = PatternFill("solid", fgColor="E5E7EB")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = Font(bold=True, color="111827")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    if rows:
        for row_idx, row in enumerate(rows, start=start_row + 1):
            for col_idx, value in enumerate(row, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value).alignment = Alignment(
                    wrap_text=True,
                    vertical="top",
                )
        return start_row + len(rows) + 3

    ws.cell(row=start_row + 1, column=1, value="No records for this period.")
    return start_row + 4


def build_excel_report(
    report: dict[str, Any], window: DateWindow, ytd_report: dict[str, Any]
) -> tuple[str, bytes]:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl is required to generate Excel reports. "
            "Install it with: pip install openpyxl"
        ) from exc

    period_label = (
        window.start_day.isoformat()
        if window.start_day == window.end_day
        else f"{window.start_day.isoformat()}_to_{window.end_day.isoformat()}"
    )
    workbook = Workbook()

    mobile_headers, mobile_rows = build_mobile_feature_usage_sheet(report)
    ws_mobile = workbook.active
    ws_mobile.title = "Mobile Feature Usage"
    ws_mobile.cell(row=1, column=1, value="Mobile Number Wise Feature Usage Counts")
    ws_mobile.cell(row=2, column=1, value=f"Period: {period_label}")
    ws_mobile.cell(row=3, column=1, value=f"Timezone: {window.timezone_name}")
    ws_mobile["A1"].font = Font(bold=True, size=14)
    ws_mobile["A2"].font = Font(color="4B5563")
    ws_mobile["A3"].font = Font(color="4B5563")
    _append_excel_table(ws_mobile, "Mobile Feature Usage", mobile_headers, mobile_rows, 5)
    ws_mobile.freeze_panes = "B7"
    _set_column_widths(ws_mobile)

    ws_summary = workbook.create_sheet("Report Summary")
    ws_summary.cell(row=1, column=1, value="Sintex Agent KPI Report")
    ws_summary.cell(row=2, column=1, value=f"Period: {period_label}")
    ws_summary.cell(row=3, column=1, value=f"Timezone: {window.timezone_name}")
    ws_summary["A1"].font = Font(bold=True, size=14)
    ws_summary["A2"].font = Font(color="4B5563")
    ws_summary["A3"].font = Font(color="4B5563")

    next_row = 5
    for title, headers, rows in build_report_summary_tables(report, window):
        next_row = _append_excel_table(ws_summary, title, headers, rows, next_row)
    _set_column_widths(ws_summary)

    ytd_window = build_ytd_window(window)
    ws_ytd = workbook.create_sheet("YTD KPI Report")
    ws_ytd.cell(row=1, column=1, value="Sintex Agent Year-to-Date KPI Report")
    ws_ytd.cell(
        row=1, column=3,
        value=f"Period: {ytd_window.start_day.isoformat()} to {ytd_window.end_day.isoformat()}",
    )
    ws_ytd["A1"].font = Font(bold=True, size=14)
    ws_ytd["C1"].font = Font(color="4B5563")
    next_row = 3
    for title, headers, rows in build_ytd_summary_tables(ytd_report, ytd_window):
        header_row = next_row + 1
        next_row = _append_excel_table(ws_ytd, title, headers, rows, next_row)
        if title == "YTD KPI Summary":
            for row_number in range(header_row + 1, header_row + len(rows) + 1):
                value_cell = ws_ytd.cell(row=row_number, column=2)
                value_cell.number_format = "#,##0.00" if isinstance(value_cell.value, float) else "#,##0"
                ws_ytd.row_dimensions[row_number].height = 45
        elif title == "Monthly Activity":
            for row_number in range(header_row + 1, header_row + len(rows) + 1):
                ws_ytd.cell(row=row_number, column=1).number_format = "mmm yyyy"
        elif title == "YTD Mobile Analytics":
            for row_number in range(header_row + 1, header_row + len(rows) + 1):
                mobile_cell = ws_ytd.cell(row=row_number, column=1)
                mobile_cell.data_type = "s"
                mobile_cell.number_format = "@"
                for column in (2, 3, 4):
                    ws_ytd.cell(row=row_number, column=column).number_format = "dd mmm yyyy"
                for column in range(5, 12):
                    ws_ytd.cell(row=row_number, column=column).number_format = "#,##0"
        ws_ytd.row_dimensions[header_row].height = 45
    ws_ytd.freeze_panes = "B5"
    _set_column_widths(ws_ytd)
    ws_ytd.column_dimensions["A"].width = 43
    ws_ytd.column_dimensions["B"].width = 24
    ws_ytd.column_dimensions["C"].width = 80
    ws_ytd.merge_cells("A1:B1")
    ws_ytd.row_dimensions[1].height = 30

    for ws in workbook.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="F3F4F6")

    output = BytesIO()
    workbook.save(output)
    filename = f"sintex_agent_kpi_report_{period_label}.xlsx"
    return filename, output.getvalue()


def build_report_html(report: dict[str, Any], window: DateWindow) -> str:
    totals = build_top_kpi_totals(report, window)

    top_kpi_rows = [
        ["Total Unique Users", fmt_int(totals["unique_users"])],
        ["New Users Came", fmt_int(totals["new_users"])],
        ["Total User Messages", fmt_int(totals["user_messages"])],
        ["Total Agent Messages", fmt_int(totals["assistant_messages"])],
    ]

    period_label = (
        window.start_day.isoformat()
        if window.start_day == window.end_day
        else f"{window.start_day.isoformat()} to {window.end_day.isoformat()}"
    )
    title_label = f"Sintex Agent KPI Report for {period_label}"

    return f"""
<!doctype html>
<html>
<head>
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  {RESPONSIVE_EMAIL_CSS}
</head>
<body style="margin:0;background:#eef2f7;font-family:Arial,Helvetica,sans-serif;color:#111827;-webkit-text-size-adjust:100%;">
  <div class="page" style="width:100%;max-width:980px;margin:0 auto;padding:20px;box-sizing:border-box;">
    <div class="hero" style="background:#111827;color:#ffffff;border-radius:8px;padding:22px 24px;">
      <div class="hero-copy" style="font-size:13px;color:#d1d5db;line-height:1.4;">{escape(window.timezone_name)}</div>
      <div class="hero-title" style="font-size:26px;font-weight:700;line-height:1.2;margin-top:4px;">{escape(title_label)}</div>
    </div>

    <section class="section" style="background:#ffffff;border:1px solid #dfe5ee;border-radius:8px;padding:18px;margin-top:14px;">
      <h2 style="font-size:17px;line-height:1.3;margin:0 0 10px;">Top KPI Summary</h2>
      {render_table(
        [
            ("KPI", "left"),
            ("Value", "right"),
        ],
        top_kpi_rows,
      )}
    </section>
  </div>
</body>
</html>
"""


def build_report_text(report: dict[str, Any], window: DateWindow) -> str:
    totals = build_top_kpi_totals(report, window)

    lines = [
        f"Sintex Agent KPI Report for {window.start_day} to {window.end_day} ({window.timezone_name})",
        "",
        "Top KPI Summary:",
        f"- Total Unique Users: {fmt_int(totals['unique_users'])}",
        f"- New Users Came: {fmt_int(totals['new_users'])}",
        f"- Total User Messages: {fmt_int(totals['user_messages'])}",
        f"- Total Agent Messages: {fmt_int(totals['assistant_messages'])}",
    ]
    return "\n".join(lines)


def send_email(
    *,
    recipients: list[str],
    subject: str,
    html_body: str,
    text_body: str,
    attachment_filename: str | None = None,
    attachment_bytes: bytes | None = None,
) -> None:
    smtp_host = os.getenv("SMTP_HOST", "smtp.office365.com")
    smtp_port = int(os.getenv("SMTP_PORT", "587"))
    smtp_user = os.getenv("SMTP_USER", "")
    smtp_pass = os.getenv("SMTP_PASS", "")
    sender = os.getenv("KPI_REPORT_FROM") or smtp_user

    if not recipients:
        raise RuntimeError(
            "No recipients configured. Set KPI_REPORT_TO or pass --to."
        )
    if not sender:
        raise RuntimeError("SMTP sender missing. Set SMTP_USER or KPI_REPORT_FROM.")
    if not smtp_user or not smtp_pass:
        raise RuntimeError("SMTP credentials missing. Set SMTP_USER and SMTP_PASS.")

    message = EmailMessage()
    message["Subject"] = subject
    message["From"] = sender
    message["To"] = ", ".join(recipients)
    message.set_content(text_body)
    message.add_alternative(html_body, subtype="html")
    if attachment_filename and attachment_bytes:
        message.add_attachment(
            attachment_bytes,
            maintype="application",
            subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            filename=attachment_filename,
        )

    try:
        logger.info(
            "Connecting to SMTP host=%s port=%s sender=%s recipients=%s",
            smtp_host,
            smtp_port,
            sender,
            len(recipients),
        )
        with smtplib.SMTP(smtp_host, smtp_port) as smtp:
            smtp.starttls()
            smtp.login(smtp_user, smtp_pass)
            smtp.send_message(message)
    except smtplib.SMTPRecipientsRefused as exc:
        refused = []
        for address, (code, reason) in exc.recipients.items():
            if isinstance(reason, bytes):
                reason = reason.decode("utf-8", errors="replace")
            refused.append(f"{address} ({code}: {reason})")
        hint = ""
        if any(address.lower().endswith("@example.com") for address in exc.recipients):
            hint = " Use a real recipient address; example.com is reserved for tests."
        raise RuntimeError(
            f"SMTP refused recipient(s): {', '.join(refused)}.{hint}"
        ) from exc
    except smtplib.SMTPAuthenticationError as exc:
        raise RuntimeError(
            "SMTP authentication failed. Check SMTP_USER and SMTP_PASS."
        ) from exc
    except smtplib.SMTPException as exc:
        raise RuntimeError(f"SMTP send failed: {exc}") from exc


async def async_main() -> None:
    args = parse_args()
    setup_logging(args.log_file)
    window = build_window(args)
    logger.info(
        "KPI report started | dry_run=%s | start_day=%s | end_day=%s | timezone=%s",
        args.dry_run,
        window.start_day,
        window.end_day,
        window.timezone_name,
    )

    database_url = os.getenv("DATABASE_URL")
    if not database_url:
        raise RuntimeError("DATABASE_URL is required to build the KPI report.")

    recipients = (
        split_csv(args.to)
        or split_csv(os.getenv("KPI_REPORT_TO"))
        or split_csv(os.getenv("KPI_EMAIL_TO"))
    )
    logger.info("Recipients resolved | count=%s", len(recipients))

    logger.info("Loading report data from database")
    report = await load_report_data(database_url, window)
    logger.info(
        "Report data loaded | daily_rows=%s | new_user_rows=%s | "
        "mobile_message_rows=%s | api_rows=%s | feature_usage_rows=%s | "
        "mobile_feature_usage_rows=%s",
        len(report["daily_rows"]),
        len(report["new_user_rows"]),
        len(report["mobile_message_rows"]),
        len(report["api_rows"]),
        len(report["feature_usage_rows"]),
        len(report["mobile_feature_usage_rows"]),
    )
    html_body = build_report_html(report, window)
    text_body = build_report_text(report, window)
    ytd_window = build_ytd_window(window)
    logger.info(
        "Loading historical YTD data | start_day=%s | end_day=%s | sources=conversations,users,messages",
        ytd_window.start_day,
        ytd_window.end_day,
    )
    ytd_report = await load_ytd_report_data(database_url, ytd_window)
    excel_filename, excel_bytes = build_excel_report(report, window, ytd_report)
    logger.info(
        "Excel report generated | filename=%s | size_bytes=%s",
        excel_filename,
        len(excel_bytes),
    )

    if args.excel_output:
        output_path = Path(args.excel_output).expanduser()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_bytes(excel_bytes)
        logger.info("Excel report saved | path=%s", output_path)

    subject_date = (
        window.start_day.isoformat()
        if window.start_day == window.end_day
        else f"{window.start_day.isoformat()} to {window.end_day.isoformat()}"
    )
    subject = f"{args.subject_prefix} - {subject_date}"
    logger.info("Report subject prepared | subject=%s", subject)

    if args.dry_run:
        print(text_body)
        print("")
        print(f"Recipients: {', '.join(recipients) if recipients else '(none)'}")
        print(f"Subject: {subject}")
        print(f"Excel attachment: {excel_filename} ({len(excel_bytes)} bytes)")
        logger.info("KPI report dry run completed")
        return

    logger.info("Sending KPI report email")
    send_email(
        recipients=recipients,
        subject=subject,
        html_body=html_body,
        text_body=text_body,
        attachment_filename=excel_filename,
        attachment_bytes=excel_bytes,
    )
    logger.info("KPI report sent successfully")
    print(f"KPI report sent to {', '.join(recipients)}")


def main() -> None:
    try:
        asyncio.run(async_main())
    except Exception as exc:
        logger.exception("KPI report failed")
        print(f"Error: {exc}", file=sys.stderr)
        raise SystemExit(1) from None


if __name__ == "__main__":
    main()
