"""Feature-result analytics regression tests.

Run: env/Scripts/python.exe -m unittest test_feature_usage -v
"""
import argparse
import importlib
import json
import logging
import sys
import types
import unittest
from datetime import date, datetime, timezone
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch

from sqlalchemy.orm import DeclarativeBase


# Import action modules without creating a real database or LLM client.
fake_db = types.ModuleType("app.database")


class Base(DeclarativeBase):
    pass


fake_db.Base = Base
fake_db.AsyncSessionLocal = None
fake_db.get_db = lambda: None
fake_chatbot = types.ModuleType("app.chatbot")
fake_chatbot.chatbot_orchestrator = MagicMock()
original_modules = {
    name: sys.modules.get(name) for name in ("app.database", "app.chatbot")
}
sys.modules.update({"app.database": fake_db, "app.chatbot": fake_chatbot})
try:
    tracker = importlib.import_module("app.analytics.feature_usage_tracker")
    events = importlib.import_module("app.analytics.feature_events")
    observability = importlib.import_module("app.utils.observability")
    rag = importlib.import_module("app.tools.sintex_rag_tool")
    qr = importlib.import_module("app.api.routes.qr_capture")
finally:
    for name, original in original_modules.items():
        if original is None:
            sys.modules.pop(name, None)
        else:
            sys.modules[name] = original

reporting = importlib.import_module("scripts.send_daily_kpi_report")


def response(body, status=200):
    result = MagicMock(status_code=status, text=json.dumps(body))
    result.json.return_value = body
    result.raise_for_status.side_effect = None
    return result


class FeatureUsageTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        logging.disable(logging.CRITICAL)
        self.addCleanup(logging.disable, logging.NOTSET)

    def test_single_date_and_inclusive_date_range_windows(self):
        single = argparse.Namespace(
            date="2026-09-08",
            from_date=None,
            to_date=None,
            yesterday=False,
            days=None,
            timezone="Asia/Kolkata",
        )
        single_window = reporting.build_window(single)
        self.assertEqual(single_window.start_day, date(2026, 9, 8))
        self.assertEqual(single_window.end_day, date(2026, 9, 8))

        date_range = argparse.Namespace(
            date=None,
            from_date="2026-09-01",
            to_date="2026-09-08",
            yesterday=False,
            days=None,
            timezone="Asia/Kolkata",
        )
        range_window = reporting.build_window(date_range)
        self.assertEqual(range_window.start_day, date(2026, 9, 1))
        self.assertEqual(range_window.end_day, date(2026, 9, 8))
        self.assertEqual(reporting.day_range(range_window)[-1], date(2026, 9, 8))

    def test_invalid_or_conflicting_date_ranges_are_rejected(self):
        base = {
            "date": None,
            "from_date": None,
            "to_date": None,
            "yesterday": False,
            "days": None,
            "timezone": "Asia/Kolkata",
        }
        invalid_cases = [
            {**base, "from_date": "2026-09-01"},
            {
                **base,
                "from_date": "2026-09-08",
                "to_date": "2026-09-01",
            },
            {
                **base,
                "date": "2026-09-08",
                "from_date": "2026-09-01",
                "to_date": "2026-09-08",
            },
            {
                **base,
                "from_date": "2026-09-01",
                "to_date": "2026-09-08",
                "days": 3,
            },
        ]
        for values in invalid_cases:
            with self.subTest(values=values), self.assertRaises(ValueError):
                reporting.build_window(argparse.Namespace(**values))

    def test_range_top_summary_counts_users_once_across_period(self):
        window = reporting.DateWindow(
            date(2026, 9, 1),
            date(2026, 9, 2),
            "UTC",
            datetime(2026, 9, 1, tzinfo=timezone.utc),
            datetime(2026, 9, 3, tzinfo=timezone.utc),
        )
        data = {
            "daily_rows": [
                {
                    "day": date(2026, 9, 1),
                    "unique_users": 1,
                    "user_messages": 2,
                    "assistant_messages": 2,
                },
                {
                    "day": date(2026, 9, 2),
                    "unique_users": 1,
                    "user_messages": 3,
                    "assistant_messages": 3,
                },
            ],
            "new_user_rows": [],
            "mobile_message_rows": [
                {"mobile_number": "919182619119", "total_messages": 5}
            ],
            "api_rows": [],
            "feature_usage_rows": [],
            "mobile_feature_usage_rows": [],
        }

        totals = reporting.build_top_kpi_totals(data, window)

        self.assertEqual(totals["unique_users"], 1)
        self.assertEqual(totals["user_messages"], 5)
        self.assertEqual(totals["assistant_messages"], 5)

    async def test_menu_selections_never_record_and_results_are_labelled(self):
        db = AsyncMock()
        factory = MagicMock()
        factory.return_value.__aenter__ = AsyncMock(return_value=db)
        factory.return_value.__aexit__ = AsyncMock(return_value=False)
        with patch.object(tracker, "AsyncSessionLocal", factory):
            for message in (
                "hi",
                "1",
                "My Rewards",
                "Product Help",
                "Upload Purchase",
                "Support",
            ):
                tracker.bind_feature_usage_context(
                    user_id=1,
                    conversation_id=2,
                    message_id=3,
                    mobile_number="test",
                    user_message=message,
                )
            db.execute.assert_not_awaited()

            await tracker.record_feature_result(
                "upload.create_claim", success=False, reason="Claim Points Not found."
            )
            failure_stmt = db.execute.call_args.args[0]
            failure_params = failure_stmt.compile().params
            self.assertEqual(failure_params["detection_method"], "failed_action")
            self.assertEqual(
                failure_params["classifier_reason"], "Claim Points Not found."
            )

            await tracker.record_feature_result("rewards.balance", success=True)
            success_stmt = db.execute.call_args.args[0]
            success_params = success_stmt.compile().params
            self.assertEqual(success_params["detection_method"], "completed_action")
            self.assertIn(
                "ON CONFLICT (message_id, event_key) DO NOTHING", str(success_stmt)
            )

    async def test_every_feature_api_records_success_and_body_failure(self):
        tracker.bind_feature_usage_context(
            user_id=1,
            conversation_id=2,
            message_id=3,
            mobile_number="test",
            user_message="feature request",
        )
        for api_name, event_key in events.API_FEATURE_EVENTS.items():
            for body, expected_success in (
                ({"status": "success"}, True),
                ({"status": "error", "message": "failed"}, False),
            ):
                client = AsyncMock()
                client.request.return_value = response(body)
                with self.subTest(api_name=api_name, body=body), patch.object(
                    observability, "record_external_api_call", AsyncMock()
                ), patch.object(
                    observability, "record_feature_result", AsyncMock()
                ) as record:
                    await observability.timed_httpx_request(
                        client,
                        "POST",
                        "https://example.invalid/test",
                        api_name=api_name,
                    )
                    self.assertEqual(record.await_count, 1)
                    self.assertEqual(record.call_args.args[0], event_key)
                    self.assertEqual(
                        record.call_args.kwargs["success"], expected_success
                    )

    async def test_api_exception_records_failed_feature_attempt(self):
        client = AsyncMock()
        client.request.side_effect = TimeoutError("offline")
        with patch.object(
            observability, "record_external_api_call", AsyncMock()
        ), patch.object(
            observability, "record_feature_result", AsyncMock()
        ) as record:
            with self.assertRaises(TimeoutError):
                await observability.timed_httpx_request(
                    client,
                    "POST",
                    "https://example.invalid/test",
                    api_name="evolve.claim_create",
                )
            record.assert_awaited_once()
            self.assertFalse(record.call_args.kwargs["success"])

    async def test_rag_records_success_and_all_post_request_failures(self):
        cases = [
            ({"chunks": []}, 200, True),
            ({"chunks": [{"text": "1000 litres", "source": "test"}]}, 200, True),
            ({"status": "error", "message": "failed"}, 200, False),
            ({}, 500, False),
            ({"chunks": "invalid"}, 200, False),
        ]
        for body, status, expected_success in cases:
            client = AsyncMock()
            client.post.return_value = response(body, status)
            factory = MagicMock()
            factory.return_value.__aenter__ = AsyncMock(return_value=client)
            factory.return_value.__aexit__ = AsyncMock(return_value=False)
            with self.subTest(body=body, status=status), patch.object(
                rag, "RAG_API_URL", "https://example.invalid"
            ), patch.object(rag.httpx, "AsyncClient", factory), patch.object(
                rag, "record_feature_result", AsyncMock()
            ) as record:
                await rag.sintex_rag_tool.coroutine("Tank capacity?")
                record.assert_awaited_once()
                self.assertEqual(
                    record.call_args.kwargs["success"], expected_success
                )

        with patch.object(rag, "record_feature_result", AsyncMock()) as record:
            await rag.sintex_rag_tool.coroutine("product help")
            record.assert_not_awaited()

    async def test_qr_records_success_and_failed_response(self):
        user = types.SimpleNamespace(id=7, mobile_number="9999999999")
        conversation = types.SimpleNamespace(id=8, language="english")
        request = qr.QRScanRequest(
            user_id_evolve=1,
            qrcode="test",
            whatsapp_number=user.mobile_number,
        )
        for body, expected_success in (
            ({"status": "success", "message": "done"}, True),
            ({"status": "error", "message": "invalid QR"}, False),
        ):
            with self.subTest(body=body), patch.object(
                qr.crud, "get_user_by_mobile", AsyncMock(return_value=user)
            ), patch.object(
                qr.crud,
                "get_or_create_conversation",
                AsyncMock(return_value=conversation),
            ), patch.object(
                qr.crud,
                "save_message",
                AsyncMock(return_value=types.SimpleNamespace(id=9)),
            ), patch.object(
                qr, "timed_requests_request", AsyncMock(return_value=response(body))
            ), patch.object(
                qr, "translate_message", AsyncMock(return_value="result")
            ), patch.object(
                qr, "send_wati_text_message", AsyncMock()
            ), patch.object(
                qr, "record_feature_result", AsyncMock()
            ) as record:
                await qr.qr_callback(request, db=AsyncMock())
                record.assert_awaited_once()
                self.assertEqual(
                    record.call_args.kwargs["success"], expected_success
                )
                self.assertEqual(record.call_args.kwargs["context"]["message_id"], 9)

    async def test_qr_exception_is_also_recorded_as_failure(self):
        user = types.SimpleNamespace(id=7, mobile_number="9999999999")
        conversation = types.SimpleNamespace(id=8, language="english")
        request = qr.QRScanRequest(
            user_id_evolve=1,
            qrcode="test",
            whatsapp_number=user.mobile_number,
        )
        with patch.object(
            qr.crud, "get_user_by_mobile", AsyncMock(return_value=user)
        ), patch.object(
            qr.crud,
            "get_or_create_conversation",
            AsyncMock(return_value=conversation),
        ), patch.object(
            qr.crud,
            "save_message",
            AsyncMock(return_value=types.SimpleNamespace(id=10)),
        ), patch.object(
            qr, "timed_requests_request", AsyncMock(side_effect=TimeoutError("offline"))
        ), patch.object(
            qr, "record_feature_result", AsyncMock()
        ) as record:
            with self.assertRaises(qr.HTTPException):
                await qr.qr_callback(request, db=AsyncMock())
            record.assert_awaited_once()
            self.assertFalse(record.call_args.kwargs["success"])
            self.assertEqual(record.call_args.kwargs["context"]["message_id"], 10)

    async def test_report_counts_successful_and_failed_attempts_as_usage(self):
        conn = AsyncMock()
        engine = MagicMock()
        engine.connect.return_value.__aenter__ = AsyncMock(return_value=conn)
        engine.connect.return_value.__aexit__ = AsyncMock(return_value=False)
        engine.dispose = AsyncMock()
        window = reporting.DateWindow(
            date(2026, 9, 6),
            date(2026, 9, 6),
            "UTC",
            datetime(2026, 9, 6, tzinfo=timezone.utc),
            datetime(2026, 9, 7, tzinfo=timezone.utc),
        )
        with patch.object(
            reporting, "create_async_engine", return_value=engine
        ), patch.object(
            reporting, "table_exists", AsyncMock(return_value=True)
        ), patch.object(
            reporting, "table_columns", AsyncMock(return_value=set())
        ), patch.object(
            reporting, "fetch_all", AsyncMock(return_value=[])
        ) as fetch:
            await reporting.load_report_data("postgresql+asyncpg://test", window)
            feature_calls = [
                call
                for call in fetch.call_args_list
                if "FROM chatbot_feature_usage_events" in call.args[1]
            ]
            self.assertEqual(len(feature_calls), 2)
            for call in feature_calls:
                self.assertIn("detection_method = ANY(:result_methods)", call.args[1])
                self.assertIn("COUNT(*) AS usage_count", call.args[1])
                self.assertNotIn("AS success_count", call.args[1])
                self.assertNotIn("AS failure_count", call.args[1])
                self.assertEqual(
                    set(call.args[2]["result_methods"]),
                    {"completed_action", "failed_action"},
                )

        data = {
            key: []
            for key in (
                "daily_rows",
                "new_user_rows",
                "mobile_message_rows",
                "api_rows",
                "feature_usage_rows",
                "mobile_feature_usage_rows",
            )
        }
        data["mobile_message_rows"] = [
            {"mobile_number": "test", "total_messages": 2}
        ]
        data["mobile_feature_usage_rows"] = [
            {
                "mobile_number": "test",
                "event_name": "Upload Purchase: Claim Created",
                "usage_count": 2,
            }
        ]
        data["feature_usage_rows"] = [
            {
                "day": date(2026, 9, 6),
                "event_name": "Upload Purchase: Claim Created",
                "flow_name": "Upload Purchase",
                "unique_users": 1,
                "usage_count": 2,
            }
        ]
        headers, rows = reporting.build_mobile_feature_usage_sheet(data)
        self.assertIn("Upload Purchase: Claim Created", headers)
        self.assertNotIn("Upload Purchase: Claim Created - Attempts", headers)
        claim_index = headers.index("Upload Purchase: Claim Created")
        self.assertEqual(rows[0][claim_index], 2)

        from openpyxl import load_workbook

        _, content = reporting.build_excel_report(data, window, {
            "activity_rows": [], "new_user_rows": [],
            "conversation_rows": [], "engagement": {}, "mobile_user_rows": [],
        })
        workbook = load_workbook(BytesIO(content))
        self.assertEqual(
            workbook.sheetnames,
            ["Mobile Feature Usage", "Report Summary", "YTD KPI Report"],
        )
        summary_tables = reporting.build_report_summary_tables(data, window)
        summary_titles = [title for title, _, _ in summary_tables]
        self.assertEqual(
            summary_titles,
            [
                "Top KPI Summary",
                "Feature Usage",
                "Flow Completion KPIs",
                "Evolve API Health",
            ],
        )
        self.assertNotIn("KPI Names", summary_titles)
        self.assertNotIn("Daily KPI Summary", summary_titles)
        self.assertNotIn("Flow Details Summary", summary_titles)
        top_summary = summary_tables[0][2]
        self.assertEqual(
            [row[0] for row in top_summary],
            [
                "Total Unique Users",
                "New Users Came",
                "Total User Messages",
                "Total Agent Messages",
            ],
        )
        html = reporting.build_report_html(data, window)
        self.assertIn("Top KPI Summary", html)
        self.assertIn("Total Unique Users", html)
        self.assertIn("Total Agent Messages", html)
        self.assertNotIn("Feature Usage", html)
        self.assertNotIn("Flow Completion KPIs", html)
        self.assertNotIn("Evolve API Health", html)
        self.assertNotIn("KPI Names", html)
        self.assertNotIn("Daily KPI Summary", html)
        self.assertNotIn("Flow Details Summary", html)
        text_report = reporting.build_report_text(data, window)
        self.assertIn("Top KPI Summary", text_report)
        self.assertIn("Total Unique Users", text_report)
        self.assertIn("Total Agent Messages", text_report)
        self.assertNotIn("Feature Usage", text_report)
        self.assertNotIn("Flow Completion KPIs", text_report)
        self.assertNotIn("Evolve API Health", text_report)
        self.assertNotIn("KPI Names", text_report)
        self.assertNotIn("Daily KPI Summary", text_report)
        self.assertNotIn("Flow Details Summary", text_report)


if __name__ == "__main__":
    unittest.main()
